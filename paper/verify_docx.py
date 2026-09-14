# -*- coding: utf-8 -*-
import io, sys, docx
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx.oxml.ns import qn

d = docx.Document('paper.docx')

print('=== 表格字号与列数')
for i, t in enumerate(d.tables):
    sizes = set()
    for row in t.rows[:2]:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    if r.font.size: sizes.add(r.font.size.pt)
    print(f' 表{i}: cols={len(t.columns)} rows={len(t.rows)} font_pt={sorted(sizes)}')

print('=== 图（含绘图段落数）')
body = d.element.xpath('//w:body')[0]
dc = 0
for p in body.findall(qn('w:p')):
    if p.find('.//' + qn('w:drawing')) is not None:
        dc += 1
print(' 绘图段落数:', dc)

print('=== 公式（oMath）个数')
print(' ', d.element.xml.count('<m:oMath'))

print('=== EN 摘要/关键词是否就位')
for p in d.paragraphs:
    if p.text.startswith('Title (EN)') or p.text.startswith('Abstract (EN)') or p.text.startswith('Keywords (EN)'):
        print(' ', p.style.name, '|', p.text[:45], '…')

print('=== 页面几何')
s = d.sections[0]
print(f" page {s.page_width.cm:.1f}x{s.page_height.cm:.1f}cm, T{s.top_margin.cm:.2f} B{s.bottom_margin.cm:.2f} L{s.left_margin.cm:.2f} R{s.right_margin.cm:.2f}")

print('=== 关键数据抽查'
      '：跨分布表 ktbd')
t6 = d.tables[6]
for r in t6.rows:
    if 'ktbd' in r.cells[0].text:
        print('  ktbd 知识点数 =', r.cells[2].text.strip(), '| AUC =', r.cells[3].text.strip())